import re
import unicodedata
from io import BytesIO
from pathlib import Path
import os
from province_codes import issuer_province
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font

FIELDS = ['Mã thẻ', 'Họ tên', 'Ngày sinh', 'Giới tính', 'Tỉnh/Thành phố', 'Nơi KCBBD', 'Hạn thẻ']

def norm(s):
    return ' '.join(unicodedata.normalize('NFC', str(s)).split()).casefold()

def input_error(text):
    # Chỉ nhận thông báo lỗi rõ ràng; không suy từ trang trống hoặc lỗi mạng.
    match = re.search(r'(mã xác nhận|captcha|ngày(?:/năm)? sinh|năm sinh|họ tên|mã số bhxh(?:/thẻ bhyt)?|mã thẻ|số cccd|cccd|thông tin nhập|dữ liệu nhập)\s+(?:không hợp lệ|không đúng|không chính xác|sai)(?!\w)', norm(text))
    return match.group(0) if match else ''

def input_result(expected, status, reason):
    result = {field: '' for field in FIELDS}
    result.update({'Trạng thái': status, 'Mã nhập': expected['Mã thẻ'],
                   'Họ tên': expected['Họ tên'], 'Ngày sinh': expected['Ngày sinh'],
                   'Nguồn thông tin': 'Họ tên/ngày sinh từ Excel nhập', 'Ghi chú': reason})
    return result

def province(address):
    # Chỉ lấy đơn vị được ghi rõ ở cuối địa chỉ, giữ tên theo thông báo.
    matches = re.findall(r'(?:^|[,;]\s*)((?:Tỉnh|Thành phố|TP\.?)\s+[^,;()]+)$', address, flags=re.I)
    if not matches:
        # Có trang không dùng dấu phẩy trước nhãn tỉnh; vẫn cần nhãn rõ ràng.
        matches = re.findall(r'\b((?:Tỉnh|Thành phố|TP\.?)\s+[^,;()]+)$', address, flags=re.I)
    return matches[0].strip() if len(matches) == 1 and '*' not in matches[0] else ''

def parse_result(text, expected):
    text = ' '.join(unicodedata.normalize('NFC', text).split())
    error = input_error(text)
    if error:
        return input_result(expected, 'Lỗi', error)
    active = 'thẻ bhyt còn giá trị sử dụng!' in norm(text)
    expired = 'thẻ bhyt hết hạn!' in norm(text)
    absent = bool(re.search(r'chưa có thông tin thẻ bhyt(?: này)? trong cơ sở dữ liệu', norm(text)))
    if absent and not active and not expired:
        return input_result(expected, 'Không có', 'Không có thông tin thẻ trong cơ sở dữ liệu được tra cứu')
    if absent:
        raise ValueError('Thông báo trạng thái mâu thuẫn; không lưu.')
    if active == expired:
        raise ValueError('Chưa có một thông báo trạng thái rõ ràng duy nhất.')
    patterns = [
        r'Mã thẻ\s*:\s*([A-Za-z0-9]+)\s*,\s*Họ tên',
        r'Họ tên\s*:\s*(.+?)\s*,\s*Ngày sinh',
        r'Ngày sinh\s*:\s*(.+?)\s*,\s*Giới tính',
        r'Giới tính\s*:\s*(.+?)\s*!?\s*\(',
        r'(?:Địa\s+chỉ|Đ\s*\.?\s*C\.?)\s*:\s*(.+?)(?=\s*[,;]?\s*Nơi\s+KCB\s*B[ĐD]\s*:)',
        r'Nơi\s+KCB\s*B[ĐD]\s*:\s*(.+?)(?=\s*[,;]?\s*Hạn\s+thẻ\s*:)',
        r'Hạn thẻ\s*:\s*(\d{2}/\d{2}/\d{4}\s*-\s*\d{2}/\d{2}/\d{4})',
    ]
    result = {}
    for field, pattern in zip(FIELDS, patterns):
        matches = re.findall(pattern, text, flags=re.I)
        if field == 'Tỉnh/Thành phố':
            labels = re.findall(r'(?:Địa\s+chỉ|Đ\s*\.?\s*C\.?)\s*:', text, flags=re.I)
            address = matches[0].strip(' ,;') if len(matches) == 1 and len(labels) == 1 else ''
            if not labels:
                regions = re.findall(r'Giới tính\s*:[^(]+\((.+?)(?=Nơi\s+KCB\s*B[ĐD]\s*:)', text, flags=re.I)
                if len(regions) == 1:
                    address = regions[0].strip(' ,;')
            result[field] = province(address)
            continue
        if len(matches) != 1 or not matches[0].strip():
            raise ValueError('Thông báo thiếu hoặc trùng trường: ' + field)
        result[field] = matches[0].strip().rstrip('!')
        if field == 'Nơi KCBBD':
            result[field] = result[field].strip(' ,;')
            if not result[field]:
                raise ValueError('Thông báo thiếu hoặc trùng trường: ' + field)
    result['Mã nhập'] = expected['Mã thẻ']
    result['Nguồn thông tin'] = 'Thông báo BHXH'
    if norm(result['Họ tên']) != norm(expected['Họ tên']):
        return input_result(expected, 'Lỗi', 'Họ tên kết quả không khớp bản ghi nhập')
    birth = expected['Ngày sinh'].strip()
    returned_birth = result['Ngày sinh']
    if re.fullmatch(r'\d{4}', birth):
        birth_ok = returned_birth[-4:] == birth
    else:
        birth_ok = returned_birth == birth or (returned_birth.startswith('**/**/') and returned_birth[-4:] == birth[-4:])
    if not birth_ok:
        return input_result(expected, 'Lỗi', 'Ngày/năm sinh kết quả không khớp bản ghi nhập')
    result['Trạng thái'] = 'Còn giá trị sử dụng' if active else 'Hết hạn'
    result['Mã tỉnh phát hành'], result['Tỉnh phát hành thẻ'] = issuer_province(result['Mã thẻ'])
    result['Nguồn tỉnh phát hành'] = 'Mã thẻ 15 ký tự / bảng mã lịch sử' if result['Tỉnh phát hành thẻ'] else ''
    return result

def excel_bytes(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'BHYT'
    columns = ['STT nguồn', 'Mã nhập'] + FIELDS + ['Mã tỉnh phát hành', 'Tỉnh phát hành thẻ', 'Nguồn tỉnh phát hành', 'Trạng thái', 'Đọc lúc', 'Nguồn thông tin', 'Ghi chú']
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([str(row.get(c, '')) for c in columns])
        color = {'Còn giá trị sử dụng': 'C6EFCE', 'Hết hạn': 'FFC7CE', 'Lỗi': 'FCE4D6', 'Bỏ qua': 'E7E6E6'}.get(row.get('Trạng thái'), 'FFF2CC')
        for cell in ws[ws.max_row]:
            cell.data_type = 's'  # Không thực thi công thức từ dữ liệu bên ngoài.
            cell.fill = PatternFill('solid', fgColor=color)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(55, max(15, max(len(str(c.value or '')) for c in col) + 2))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()

def save_atomic(path, rows):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix('.tmp')
    try:
        temp.write_bytes(excel_bytes(rows))
        os.replace(temp, path)
    finally:
        if temp.exists():
            temp.unlink()
