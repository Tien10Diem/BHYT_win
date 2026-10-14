from datetime import datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4
import re
import os
import streamlit as st
from openpyxl import load_workbook
from browser import Browser
from core import parse_result, save_atomic, excel_bytes, input_result

RESULTS_DIR = Path(os.environ.get('BHYT_DATA_DIR', str(Path(__file__).parent))) / 'results'

st.set_page_config(page_title='Tra cứu BHYT', layout='wide')
st.title('Tra cứu BHYT • OCR CAPTCHA tự động')
st.caption('Tự động điền thông tin, OCR mã CAPTCHA, đọc kết quả và chuyển tiếp. Hỗ trợ nhập tay khi OCR sai.')
ss = st.session_state
for key, value in {'rows': [], 'results': [], 'index': 0, 'browser': None, 'prepared': None, 'notice': '', 'auto_paused': False}.items():
    if key not in ss:
        ss[key] = value
if 'output' not in ss:
    ss.output = RESULTS_DIR / ('bhyt_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '_' + uuid4().hex[:8] + '.xlsx')

def prepare_current(auto_ocr=False):
    ss.prepared = None
    if ss.browser is None:
        ss.browser = Browser()
    ss.browser.prepare(ss.rows[ss.index], auto_ocr=auto_ocr)
    ss.prepared = ss.index
    ss.auto_paused = False

def commit_and_advance(result):
    result.update({'STT nguồn': ss.rows[ss.index]['STT nguồn'], 'Đọc lúc': datetime.now().astimezone().isoformat(timespec='seconds')})
    new_results = ss.results + [result]
    save_atomic(ss.output, new_results)
    ss.results = new_results
    ss.index += 1
    ss.prepared = None
    if ss.browser:
        ss.browser.ready = False
    ss.notice = 'Đã lưu: ' + result['Trạng thái'] + '.'
    if ss.index < len(ss.rows):
        use_ocr = ss.get('ocr_enabled', True)
        try:
            prepare_current(auto_ocr=use_ocr)
            if use_ocr and ss.browser and ss.browser.last_ocr_text:
                ss.notice += f' Đã điền bản ghi tiếp + OCR CAPTCHA: "{ss.browser.last_ocr_text}". Nhập CAPTCHA trong Chrome nếu cần sửa.'
            else:
                ss.notice += ' Đã điền bản ghi tiếp theo; nhập CAPTCHA trong Chrome.'
        except Exception:
            ss.notice += ' Chưa điền được dòng tiếp theo; bấm Mở / điền để thử lại.'

def read_input(data):
    # Chỉ XLSX, không chạy macro, giữ mã thẻ là chuỗi và không đọc công thức.
    wb = load_workbook(BytesIO(data), read_only=True, data_only=False)
    try:
        ws = wb.active
        if ws.max_row > 10001 or ws.max_column > 100:
            raise ValueError('Giới hạn 10.000 dòng và 100 cột.')
        it = ws.iter_rows()
        headers = [str(c.value or '').strip() for c in next(it)]
        required = ['Mã thẻ', 'Họ tên', 'Ngày sinh']
        if any(headers.count(h) != 1 for h in required):
            raise ValueError('Cần ba cột duy nhất: Mã thẻ, Họ tên, Ngày sinh.')
        positions = [headers.index(h) for h in required]
        rows = []
        for number, cells in enumerate(it, 2):
            if all(c.value is None for c in cells):
                continue
            selected = [cells[i] for i in positions]
            if any(c.data_type == 'f' or not isinstance(c.value, str) for c in selected):
                raise ValueError(f'Dòng {number}: ba trường phải là Text, không phải công thức/số/ngày Excel.')
            row = dict(zip(required, [c.value.strip() for c in selected]))
            if not re.fullmatch(r'[A-Za-z0-9]{10,15}', row['Mã thẻ']) or not row['Họ tên']:
                raise ValueError(f'Dòng {number}: mã thẻ/họ tên không hợp lệ.')
            if not re.fullmatch(r'(?:\d{4}|\d{2}/\d{2}/\d{4})', row['Ngày sinh']):
                raise ValueError(f'Dòng {number}: ngày sinh dùng dd/mm/yyyy hoặc yyyy.')
            row['STT nguồn'] = str(number)
            rows.append(row)
        if not rows:
            raise ValueError('Excel không có bản ghi.')
        return rows
    finally:
        wb.close()

left, right = st.columns([2, 1])

@st.fragment(run_every='1s')
def watch_result():
    if not ss.get('auto_enabled', True) or ss.auto_paused or ss.prepared != ss.index or not ss.browser or ss.index >= len(ss.rows):
        return
    try:
        text = ss.browser.auto_read()
        if text:
            result = parse_result(text, ss.rows[ss.index])
            commit_and_advance(result)
            st.rerun()
    except Exception:
        ss.auto_paused = True
        st.warning('Tạm dừng tự đọc vì chưa đọc/lưu được. Kiểm tra Chrome/Excel, dùng nút Đọc kết quả và chuyển tiếp hoặc mở lại bản ghi để tiếp tục tự đọc.')

with left:
    col_toggle1, col_toggle2 = st.columns(2)
    with col_toggle1:
        st.toggle('Tự đọc kết quả và chuyển tiếp', value=True, key='auto_enabled', help='Sau khi bấm Tra cứu, chờ thông báo mới ổn định 3 giây rồi lưu và chuyển dòng.')
    with col_toggle2:
        st.toggle('🔍 OCR CAPTCHA tự động', value=True, key='ocr_enabled', help='Tự động đọc và điền mã CAPTCHA khi mở bản ghi. Tắt nếu muốn nhập tay.')

    upload = st.file_uploader('Danh sách Excel (.xlsx)', type=['xlsx'], disabled=bool(ss.rows))
    if upload and not ss.rows and st.button('Nạp danh sách'):
        try:
            ss.rows = read_input(upload.getvalue())
            st.rerun()
        except Exception as exc:
            st.error(str(exc) if isinstance(exc, ValueError) else 'Không đọc được Excel. Kiểm tra định dạng XLSX.')
    st.caption('Cột bắt buộc: Mã thẻ, Họ tên, Ngày sinh. Định dạng Text; ngày sinh dd/mm/yyyy hoặc yyyy.')
    if ss.rows:
        st.progress(ss.index / len(ss.rows), text=f'Đã lưu {ss.index}/{len(ss.rows)} bản ghi')
        if ss.index < len(ss.rows):
            row = ss.rows[ss.index]
            st.write('Bản ghi hiện tại', row)

            # --- Các nút thao tác chính ---
            btn_col1, btn_col2 = st.columns(2)
            with btn_col1:
                if st.button('📝 Mở / điền bản ghi trong Chrome'):
                    try:
                        use_ocr = ss.get('ocr_enabled', True)
                        prepare_current(auto_ocr=use_ocr)
                        if use_ocr and ss.browser and ss.browser.last_ocr_text:
                            ss.notice = f'Đã điền + OCR CAPTCHA: "{ss.browser.last_ocr_text}". Kiểm tra và bấm Tra cứu trong Chrome.'
                        else:
                            ss.notice = 'Đã điền. Nhập CAPTCHA và bấm Tra cứu trực tiếp trong Chrome.'
                    except Exception:
                        ss.notice = 'Không mở/điền được Chrome. Kiểm tra Chrome, mạng và selector; bản ghi chưa được lưu.'
            with btn_col2:
                if st.button('🤖 OCR + Tra cứu tự động', disabled=ss.prepared != ss.index,
                             help='OCR mã CAPTCHA hiện tại, điền vào ô và bấm Tra cứu tự động.'):
                    if ss.browser:
                        try:
                            ocr_text, submitted = ss.browser.solve_and_submit()
                            if submitted:
                                ss.notice = f'OCR: "{ocr_text}" → Đã bấm Tra cứu. Chờ kết quả...'
                            elif ocr_text:
                                ss.notice = f'OCR: "{ocr_text}" nhưng không bấm được Tra cứu. Bấm thủ công trong Chrome.'
                            else:
                                ss.notice = ss.browser.last_ocr_status
                        except Exception as e:
                            ss.notice = f'Lỗi OCR/Tra cứu: {str(e)}'
                    else:
                        ss.notice = 'Chưa mở Chrome. Bấm "Mở / điền bản ghi" trước.'

            if st.button('📖 Đọc kết quả và chuyển tiếp', type='primary', disabled=ss.prepared != ss.index):
                try:
                    result = parse_result(ss.browser.read(), row)
                    commit_and_advance(result)
                    st.rerun()
                except ValueError as exc:
                    st.error(str(exc))
                except PermissionError:
                    st.error('Không lưu được Excel. Đóng file kết quả đang mở rồi thử lại; chưa chuyển dòng.')
                except Exception:
                    st.error('Đọc hoặc lưu thất bại. Không chuyển dòng; kiểm tra Chrome, thư mục results và thử lại.')

            skip_col1, skip_col2 = st.columns(2)
            with skip_col1:
                if st.button('⏭️ Bỏ qua dòng này', help='Không đọc hoặc tra cứu dòng hiện tại; lưu Bỏ qua và chuyển dòng.'):
                    try:
                        commit_and_advance(input_result(row, 'Bỏ qua', 'Người dùng bỏ qua, chưa tra cứu'))
                        st.rerun()
                    except PermissionError:
                        st.error('Không lưu được Excel. Đóng file kết quả đang mở rồi thử lại; chưa chuyển dòng.')
                    except Exception:
                        st.error('Không lưu được dòng Bỏ qua; chưa chuyển dòng.')
            with skip_col2:
                if st.button('🔄 OCR lại CAPTCHA', disabled=ss.prepared != ss.index,
                             help='Đọc lại mã CAPTCHA mới (khi trang đã tải lại CAPTCHA).'):
                    if ss.browser:
                        try:
                            text = ss.browser.solve_captcha()
                            if text:
                                ss.notice = f'OCR lại: "{text}". Bấm Tra cứu trong Chrome hoặc dùng nút "OCR + Tra cứu tự động".'
                            else:
                                ss.notice = ss.browser.last_ocr_status
                        except Exception as e:
                            ss.notice = f'Lỗi OCR: {str(e)}'

            # Hiển thị trạng thái OCR
            if ss.browser and ss.browser.last_ocr_status and ss.prepared == ss.index:
                st.info(f'🔍 {ss.browser.last_ocr_status}')
        else:
            st.success('Đã xử lý hết danh sách.')
            if st.button('Tra cứu lại', help='Giữ danh sách nhập, bắt đầu từ dòng đầu và tạo file kết quả mới.'):
                ss.results = []
                ss.index = 0
                ss.prepared = None
                ss.auto_paused = False
                ss.output = RESULTS_DIR / ('bhyt_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '_' + uuid4().hex[:8] + '.xlsx')
                try:
                    prepare_current(auto_ocr=ss.get('ocr_enabled', True))
                    ss.notice = 'Bắt đầu tra cứu lại từ dòng đầu; kết quả cũ vẫn được giữ trong results.'
                except Exception:
                    ss.notice = 'Đã đặt lại danh sách. Bấm Mở / điền bản ghi trong Chrome để bắt đầu.'
                st.rerun()
    watch_result()
    if ss.notice:
        st.info(ss.notice)
    if ss.results:
        st.dataframe(ss.results, hide_index=True)
        st.download_button('Tải Excel kết quả', data=excel_bytes(ss.results), file_name=ss.output.name,
                           mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        st.caption('Bản tự lưu: ' + str(ss.output))
with right:
    st.subheader('Chrome và OCR')
    st.write('🟢 **OCR tự động**: Khi bật, ứng dụng sẽ tự chụp ảnh CAPTCHA, nhận dạng text bằng AI và điền vào ô mã xác nhận.')
    st.write('🔵 **OCR + Tra cứu**: Bấm nút để OCR CAPTCHA và tự động bấm Tra cứu luôn.')
    st.write('🟡 **OCR lại**: Khi CAPTCHA bị sai, trang sẽ tải lại CAPTCHA mới. Bấm nút này để đọc lại.')
    st.divider()
    st.write('Xanh: còn giá trị sử dụng. Đỏ: hết hạn. Vàng: Không có. Cam: Lỗi. Xám: Bỏ qua.')
    st.caption('Tỉnh phát hành thẻ là cột riêng theo mã thẻ 15 ký tự và bảng mã lịch sử. Tự đọc chỉ hoạt động khi theo dõi được thao tác Tra cứu và kết quả mới.')
    st.divider()
    if st.button('🚫 Đóng Chrome'):
        if ss.browser:
            try:
                ss.browser.close()
            except Exception:
                pass
        ss.browser = None
        ss.prepared = None
        ss.notice = 'Đã đóng Chrome. Có thể mở lại bản ghi hiện tại.'
        st.rerun()
    st.caption('Chạy một người dùng trên máy Windows cá nhân. Không đưa dữ liệu lên dịch vụ công khai.')
