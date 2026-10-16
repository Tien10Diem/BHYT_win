from datetime import datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4
import re
import os
import time
from time import sleep
import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from browser import Browser
from core import parse_result, save_atomic, excel_bytes, input_result

RESULTS_DIR = Path(os.environ.get('BHYT_DATA_DIR', str(Path(__file__).parent))) / 'results'

st.set_page_config(page_title='Tra cứu BHYT', layout='wide')
st.title('Vua Tiến \n Vua đất Huế')
st.caption('chưa nghĩ ra caption')

ss = st.session_state
defaults = {
    'rows': [],
    'results': [],
    'index': 0,
    'browser': None,
    'notice': '',
    'auto_running': False,
    'delay_sec': 0.2,
    'max_retries': 3,
    'start_time': None,
    'total_elapsed': 0.0
}
for key, value in defaults.items():
    if key not in ss:
        ss[key] = value

def get_elapsed_sec():
    if getattr(ss, 'auto_running', False) and getattr(ss, 'start_time', None) is None:
        ss.start_time = time.time()

    elapsed = getattr(ss, 'total_elapsed', 0.0)
    if getattr(ss, 'auto_running', False) and getattr(ss, 'start_time', None):
        elapsed += time.time() - ss.start_time
    return elapsed

def format_time(seconds):
    seconds = max(0, int(seconds))
    h = seconds // 3600
    m = (seconds % 3600) // 60
    s = seconds % 60
    if h > 0:
        return f"{h:02d}:{m:02d}:{s:02d}"
    return f"{m:02d}:{s:02d}"

if 'output' not in ss:
    ss.output = RESULTS_DIR / ('bhyt_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '_' + uuid4().hex[:8] + '.xlsx')

def read_input(data):
    wb = load_workbook(BytesIO(data), read_only=True, data_only=False)
    try:
        ws = wb.active
        if ws.max_row > 10001 or ws.max_column > 100:
            raise ValueError('Giới hạn 10.000 dòng và 100 cột.')
        it = ws.iter_rows()
        headers = [str(c.value or '').strip() for c in next(it)]
        required = ['Mã thẻ', 'Họ tên', 'Ngày sinh']
        if any(headers.count(h) != 1 for h in required):
            raise ValueError('Cần ba cột duy nhất: Mã thẻ, Họ tên, Ngày sinh.')
        positions = [headers.index(h) for h in required]
        rows = []
        for number, cells in enumerate(it, 2):
            if all(c.value is None for c in cells):
                continue
            selected = [cells[i] for i in positions]
            if any(c.data_type == 'f' or not isinstance(c.value, str) for c in selected):
                raise ValueError(f'Dòng {number}: ba trường phải là Text, không phải công thức/số/ngày Excel.')
            row = dict(zip(required, [c.value.strip() for c in selected]))
            if not re.fullmatch(r'[A-Za-z0-9]{10,15}', row['Mã thẻ']) or not row['Họ tên']:
                raise ValueError(f'Dòng {number}: mã thẻ/họ tên không hợp lệ.')
            if not re.fullmatch(r'(?:\d{4}|\d{2}/\d{2}/\d{4})', row['Ngày sinh']):
                raise ValueError(f'Dòng {number}: ngày sinh dùng dd/mm/yyyy hoặc yyyy.')
            row['STT nguồn'] = str(number)
            rows.append(row)
        if not rows:
            raise ValueError('Excel không có bản ghi.')
        return rows
    finally:
        wb.close()

def commit_result(row, result):
    result.update({
        'STT nguồn': row['STT nguồn'],
        'Đọc lúc': datetime.now().astimezone().isoformat(timespec='seconds')
    })
    new_results = ss.results + [result]
    ss.results = new_results
    try:
        save_atomic(ss.output, new_results)
    except PermissionError:
        ss.notice = '⚠️ Không thể ghi file kết quả do file đang được mở trong Excel. Vui lòng đóng file để lưu tiếp.'
    except Exception as e:
        ss.notice = f'Lỗi lưu kết quả: {str(e)}'
    ss.index += 1

# Bố cục 2 cột: Trái (Điều khiển & Tiến độ), Phải (Bảng kết quả trực tiếp)
left, right = st.columns([1, 1.35])

with left:
    # 1. Nạp file Excel nếu chưa có dữ liệu
    if not ss.rows:
        upload = st.file_uploader('Chọn file danh sách Excel (.xlsx)', type=['xlsx'])
        if upload and st.button('📥 Nạp danh sách và chuẩn bị', type='primary'):
            try:
                ss.rows = read_input(upload.getvalue())
                ss.index = 0
                ss.results = []
                ss.start_time = None
                ss.total_elapsed = 0.0
                ss.notice = f'Đã nạp thành công {len(ss.rows)} bản ghi. Hãy bấm "Bắt đầu chạy tự động".'
                st.rerun()
            except Exception as exc:
                st.error(str(exc) if isinstance(exc, ValueError) else 'Lỗi đọc file Excel. Kiểm tra định dạng.')
        st.caption('Cột bắt buộc trong file: Mã thẻ, Họ tên, Ngày sinh (định dạng Text, ngày sinh dd/mm/yyyy hoặc yyyy).')
    else:
        # 2. Điều khiển Tự Động Toàn Bộ
        st.subheader('⚙️ Điều Khiển Tự Động')

        ctrl_col1, ctrl_col2 = st.columns([1.3, 1])
        with ctrl_col1:
            if not ss.auto_running:
                if ss.index < len(ss.rows):
                    btn_text = '▶️ TIẾP TỤC CHẠY TỰ ĐỘNG' if ss.index > 0 else '▶️ BẮT ĐẦU CHẠY TỰ ĐỘNG TOÀN BỘ'
                    if st.button(btn_text, type='primary', width='stretch'):
                        ss.auto_running = True
                        ss.start_time = time.time()
                        ss.notice = 'Đang bắt đầu chạy tự động...'
                        st.rerun()
                else:
                    st.success('🎉 Đã hoàn thành tra cứu toàn bộ danh sách!')
            else:
                if st.button('⏸️ TẠM DỪNG TỰ ĐỘNG', type='secondary', width='stretch'):
                    ss.auto_running = False
                    if getattr(ss, 'start_time', None):
                        ss.total_elapsed = getattr(ss, 'total_elapsed', 0.0) + (time.time() - ss.start_time)
                        ss.start_time = None
                    ss.notice = 'Đã tạm dừng tự động.'
                    st.rerun()

        with ctrl_col2:
            ss.delay_sec = st.slider('Nghỉ giữa mỗi case (giây)', min_value=0.0, max_value=3.0, value=float(ss.delay_sec), step=0.1)

        # Thanh tiến độ
        progress_val = min(1.0, ss.index / max(1, len(ss.rows)))
        st.progress(progress_val, text=f'Tiến độ: {ss.index} / {len(ss.rows)} bản ghi ({int(progress_val * 100)}%)')

        # Bộ đếm thời gian từ lúc bấm nút Bắt đầu đến khi hết file (tạm dừng thì dừng đếm)
        elapsed_sec = get_elapsed_sec()
        elapsed_str = format_time(elapsed_sec)
        if ss.index >= len(ss.rows) and len(ss.rows) > 0:
            st.success(f"🎉 **Tổng thời gian hoàn thành toàn bộ file:** `{elapsed_str}`")
        elif ss.auto_running:
            st.info(f"⏱️ **Thời gian chạy từ lúc bấm Bắt đầu:** `{elapsed_str}` *(Đang đếm giờ...)*")
        elif elapsed_sec > 0:
            st.warning(f"⏱️ **Thời gian chạy từ lúc bấm Bắt đầu:** `{elapsed_str}` *(Đang tạm dừng)*")
        else:
            st.caption(f"⏱️ **Thời gian chạy:** `00:00` *(Bắt đầu đếm khi bạn bấm nút chạy tự động)*")

        if ss.index < len(ss.rows):
            current_row = ss.rows[ss.index]
            with st.container(border=True):
                st.markdown(f"**Đang xử lý:** `{current_row['Họ tên']}` | Ngày sinh: `{current_row['Ngày sinh']}` | Mã: `{current_row['Mã thẻ']}` (Dòng {ss.index + 1}/{len(ss.rows)})")
                if ss.browser and ss.browser.last_ocr_status:
                    st.caption(f"🔍 Trạng thái OCR: {ss.browser.last_ocr_status}")

            # Thao tác thủ công bổ trợ (khi cần)
            with st.expander('🛠️ Thao tác thủ công bổ trợ (tùy chọn)'):
                m_col1, m_col2 = st.columns(2)
                with m_col1:
                    if st.button('⏭️ Bỏ qua dòng này', width='stretch', help='Bỏ qua người này và chuyển sang dòng tiếp theo.'):
                        commit_result(current_row, input_result(current_row, 'Bỏ qua', 'Người dùng bỏ qua, chưa tra cứu'))
                        st.rerun()
                with m_col2:
                    if st.button('🔄 Tra cứu lại từ đầu', width='stretch', help='Đặt lại từ dòng 1 và tạo file kết quả mới.'):
                        ss.results = []
                        ss.index = 0
                        ss.auto_running = False
                        ss.start_time = None
                        ss.total_elapsed = 0.0
                        ss.output = RESULTS_DIR / ('bhyt_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '_' + uuid4().hex[:8] + '.xlsx')
                        ss.notice = 'Đã đặt lại danh sách về dòng đầu.'
                        st.rerun()

        if ss.notice:
            st.info(ss.notice)

with right:
    st.subheader('📋 Bảng Theo Dõi Trực Tiếp Từng Case')

    if ss.results:
        df = pd.DataFrame(ss.results)

        # Thống kê nhanh kết quả
        counts = df['Trạng thái'].value_counts().to_dict() if 'Trạng thái' in df.columns else {}
        stat_parts = [f"**Tổng**: {len(df)}/{len(ss.rows)}"]
        if 'Còn giá trị sử dụng' in counts:
            stat_parts.append(f"🟢 **Còn hạn**: {counts['Còn giá trị sử dụng']}")
        if 'Hết hạn' in counts:
            stat_parts.append(f"🔴 **Hết hạn**: {counts['Hết hạn']}")
        if 'Lỗi' in counts:
            stat_parts.append(f"🟠 **Lỗi**: {counts['Lỗi']}")
        if 'Không có' in counts:
            stat_parts.append(f"🟡 **Không có**: {counts['Không có']}")
        if 'Bỏ qua' in counts:
            stat_parts.append(f"⚪ **Bỏ qua**: {counts['Bỏ qua']}")
        st.markdown(" | ".join(stat_parts))

        # Nút tải file Excel
        st.download_button(
            '📥 Tải file Excel kết quả hiện tại',
            data=excel_bytes(ss.results),
            file_name=ss.output.name,
            mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            type='primary',
            width='stretch'
        )

        # Hiển thị bảng kết quả với case mới nhất nằm trên cùng (tối đa 50 case gần nhất để UI mượt mà, file Excel lưu đầy đủ 100%)
        cols = [c for c in ['STT nguồn', 'Họ tên', 'Ngày sinh', 'Mã thẻ', 'Trạng thái', 'Hạn thẻ', 'Ghi chú'] if c in df.columns]
        st.dataframe(df[cols].iloc[::-1].head(50), hide_index=True, width='stretch', height=520)
        st.caption(f'📁 Hiển thị 50 case mới nhất trên bảng. Toàn bộ {len(df)}/{len(ss.rows)} kết quả được lưu tự động đầy đủ 100% tại: `{ss.output}`')
    else:
        st.info('👉 Kết quả từng case sẽ tự động cập nhật lên bảng này ngay khi tra cứu xong từng người (case mới nhất nằm trên cùng).')

    st.divider()
    if st.button('🚫 Đóng trình duyệt Chrome'):
        if ss.browser:
            try:
                ss.browser.close()
            except Exception:
                pass
        ss.browser = None
        ss.auto_running = False
        ss.notice = 'Đã đóng Chrome.'
        st.rerun()

# --- VÒNG LẶP CHẠY TỰ ĐỘNG TOÀN BỘ (ĐẶT Ở CUỐI SCRIPT ĐỂ UI LUÔN HIỂN THỊ TRƯỚC) ---
if ss.auto_running and ss.index < len(ss.rows):
    current_row = ss.rows[ss.index]
    result = None

    try:
        # Cứ mỗi 50 case mới cần làm sạch trình duyệt hoàn toàn; tránh việc đóng mở Chrome liên tục mỗi 20 case làm đứng hình 5 giây
        if ss.index > 0 and ss.index % 50 == 0 and getattr(ss, 'last_recreated_index', -1) != ss.index:
            ss.last_recreated_index = ss.index
            if ss.browser:
                try:
                    ss.browser.close()
                except Exception:
                    pass
                ss.browser = None

        if ss.browser is None:
            ss.browser = Browser()

        # Thực hiện quy trình tra cứu cho bản ghi hiện tại
        status, payload = ss.browser.process_record(current_row)

        if status == 'SUCCESS':
            ss.captcha_fail_count = 0
            result = parse_result(payload, current_row)
            commit_result(current_row, result)
            ss.notice = f"Đã lưu: {result['Trạng thái']} ({current_row['Họ tên']})"
        elif status == 'DATA_ERROR':
            ss.captcha_fail_count = 0
            result = input_result(current_row, 'Lỗi', payload)
            commit_result(current_row, result)
            ss.notice = f"Đã ghi nhận lỗi sai thông tin: {payload} ({current_row['Họ tên']})"
        else:
            # Không lấy được kết quả từ BHXH (thẻ không tồn tại/lỗi): Ghi nhận 'Không có thông tin' và NEXT ngay!
            result = input_result(current_row, 'Không có', 'Không có thông tin')
            commit_result(current_row, result)
            ss.notice = f"Đã ghi nhận: Không có thông tin ({current_row['Họ tên']}) ➔ Chuyển sang người tiếp theo"

    except Exception as exc:
        # Nếu gặp ngoại lệ bất ngờ: Ghi nhận 'Không có thông tin' và NEXT ngay, không dừng lại!
        result = input_result(current_row, 'Không có', 'Không có thông tin')
        commit_result(current_row, result)
        if ss.browser:
            try:
                ss.browser.close()
            except Exception:
                pass
            ss.browser = None
        ss.notice = f"Đã ghi nhận: Không có thông tin ({current_row['Họ tên']}) ➔ Chuyển sang người tiếp theo"

    # Nghỉ giữa mỗi case theo thiết lập kèm nhịp thở tự nhiên (chống tường lửa BHXH bóp băng thông)
    if result and ss.delay_sec > 0:
        import random
        sleep(ss.delay_sec + random.uniform(0.05, 0.25))

    # Nếu đã chạy hết danh sách
    if ss.index >= len(ss.rows):
        ss.auto_running = False
        if getattr(ss, 'start_time', None):
            ss.total_elapsed = getattr(ss, 'total_elapsed', 0.0) + (time.time() - ss.start_time)
            ss.start_time = None
        ss.notice = '🎉 Đã hoàn thành tra cứu toàn bộ danh sách!'

    # Tự động chuyển và nạp case tiếp theo
    st.rerun()
